"""Read-only evidence generator for the 0-D data × legacy Ontology audit.

It does not build a query store, Registry, Ontology, or runtime binding.  It
only profiles immutable inputs and writes reproducible provenance catalogs.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import tempfile
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import duckdb
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OFFICIAL = ROOT / "data" / "official_raw"
HOLDINGS_ZIP = ROOT / "data" / "incoming" / "holdings_20260829.zip"
LEGACY_ONTOLOGY = Path(r"C:\Users\user\asset_agent\canna\ontology")
LEGACY_MAPPING = Path(r"C:\Users\user\asset_agent\canna\docs\ontology\RDB_ONTOLOGY_MAPPING.md")

TABLES = {
    "PRBD01N001": {
        "data": "prbd01n001_data.xlsx",
        "schema": "prbd01n001_schema.xlsx",
        "source_row_grain": "(pd_no, pd_exg_mkt, info_seq)",
        "identifier_candidates": "pd_no (product candidate); composite source-row key",
        "as_of_fields": "info_base_dt, pd_std_info_update, sale_yield_base_dt, crd_grd_dt",
    },
    "PREF01N001": {
        "data": "pref01n001_data.xlsx",
        "schema": "pref01n001_schema.xlsx",
        "source_row_grain": "provided exchange-traded product/listing snapshot row",
        "identifier_candidates": "pd_itm_no, pd_itm_no_ma; pd_isin_cd/pd_ric/pd_ticker are verification candidates",
        "as_of_fields": "field-specific *_base_dt, *_upt_dt, pd_dvid_prc_base_dt, cu_upt_dt",
    },
    "PREF02N001": {
        "data": "pref02n001_data.xlsx",
        "schema": "pref02n001_schema.xlsx",
        "source_row_grain": "provided overseas exchange-traded product/listing snapshot row",
        "identifier_candidates": "pd_itm_no; pd_isin_cd/pd_lipper_id/pd_us_cik are distinct candidates",
        "as_of_fields": "field-specific du_clpr_base_dt, du_nav_base_dt, du_upt_dt, cu_upt_dt, wu_upt_dt",
    },
    "PRFD01N001": {
        "data": "prfd01n001_data.xlsx",
        "schema": "prfd01n001_schema.xlsx",
        "source_row_grain": "fund share-class row (itm_no)",
        "identifier_candidates": "itm_no, std_itm_no, ksd_itm_no, rptt_ksd_itm_no, fss_itm_no, mtco_itm_no",
        "as_of_fields": "fd_daily_bas_dt, fd_price_bas_dt, fd_last_dstb_actg_bss_dt, fd_last_dstb_actg_eot_dt",
    },
}

DECLARATION_TYPES = (
    "owl:Class",
    "owl:ObjectProperty",
    "owl:DatatypeProperty",
    "owl:AnnotationProperty",
    "fpa:MetricType",
    "sh:NodeShape",
    "fpa:IdentifierScheme",
    "fpa:ValueStatus",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def zero_like(value: object) -> bool:
    if isinstance(value, (int, float)):
        return value == 0
    if isinstance(value, str):
        return value.strip() in {"0", "0.0", "0.00", "0.000", "0.0000"}
    return False


def read_schema(path: Path) -> tuple[list[str], dict[str, dict[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headings = [str(value).strip() if value is not None else "" for value in next(rows)]
        index = {heading: position for position, heading in enumerate(headings)}
        required = {"컬럼명", "데이터타입", "Nullable", "컬럼코멘트"}
        missing = required.difference(index)
        if missing:
            raise ValueError(f"{path.name}: missing schema headings {sorted(missing)}")
        result: dict[str, dict[str, str]] = {}
        order: list[str] = []
        for row in rows:
            field = row[index["컬럼명"]]
            if field is None:
                continue
            name = str(field).strip()
            order.append(name)
            result[name] = {
                "data_type": str(row[index["데이터타입"]] or ""),
                "nullable": str(row[index["Nullable"]] or ""),
                "comment": str(row[index["컬럼코멘트"]] or ""),
            }
        return order, result
    finally:
        workbook.close()


def legacy_binding_lines(field_names: set[str]) -> dict[str, tuple[str, int, str]]:
    statuses = (
        "DIRECT",
        "OBSERVATION",
        "ROLE",
        "ENTITY_RESOLUTION",
        "RAW_CLASSIFICATION",
        "EXCLUDED_INFERENCE",
        "EXTERNAL_REQUIRED",
    )
    results: dict[str, tuple[str, int, str]] = {}
    for line_number, line in enumerate(LEGACY_MAPPING.read_text(encoding="utf-8").splitlines(), 1):
        mentioned = set(re.findall(r"`([^`]+)`", line)).intersection(field_names)
        found_statuses = [status for status in statuses if status in line]
        if not mentioned:
            continue
        for field in mentioned:
            results.setdefault(field, ("+".join(found_statuses) if found_statuses else "LEGACY_TERM_REFERENCE", line_number, line.strip()))
    return results


def official_catalog() -> tuple[list[dict[str, object]], list[dict[str, object]], set[str]]:
    all_fields: set[str] = set()
    schemas: dict[str, tuple[list[str], dict[str, dict[str, str]]]] = {}
    for table_id, metadata in TABLES.items():
        order, schema = read_schema(OFFICIAL / str(metadata["schema"]))
        schemas[table_id] = (order, schema)
        all_fields.update(order)
    legacy = legacy_binding_lines(all_fields)
    catalog: list[dict[str, object]] = []
    source_summary: list[dict[str, object]] = []

    for table_id, metadata in TABLES.items():
        data_path = OFFICIAL / str(metadata["data"])
        schema_path = OFFICIAL / str(metadata["schema"])
        order, schema = schemas[table_id]
        counts = {field: Counter() for field in order}
        workbook = load_workbook(data_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            data_headers = [str(value).strip() if value is not None else "" for value in next(rows)]
            if data_headers != order:
                raise ValueError(f"{table_id}: data header order differs from schema")
            row_count = 0
            for row in rows:
                row_count += 1
                for position, value in enumerate(row):
                    if value is None or (isinstance(value, str) and not value.strip()):
                        continue
                    field = order[position]
                    counts[field]["nonblank"] += 1
                    if zero_like(value):
                        counts[field]["zero_like"] += 1
                    else:
                        counts[field]["nonzero_nonblank"] += 1
        finally:
            workbook.close()

        source_summary.append(
            {
                "table_id": table_id,
                "data_file": metadata["data"],
                "schema_file": metadata["schema"],
                "data_sha256": sha256(data_path),
                "schema_sha256": sha256(schema_path),
                "data_rows": row_count,
                "data_columns": len(order),
                "source_row_grain": metadata["source_row_grain"],
                "identifier_candidates": metadata["identifier_candidates"],
                "as_of_fields": metadata["as_of_fields"],
                "coverage_failure_source": "official source has no collection failure table; field counts below are observations, not claim policy",
            }
        )
        for field in order:
            mapping = legacy.get(field)
            catalog.append(
                {
                    "table_id": table_id,
                    "data_file": metadata["data"],
                    "schema_file": metadata["schema"],
                    "source_row_grain": metadata["source_row_grain"],
                    "identifier_candidates": metadata["identifier_candidates"],
                    "as_of_fields": metadata["as_of_fields"],
                    "field": field,
                    **schema[field],
                    "data_rows": row_count,
                    "nonblank_rows": counts[field]["nonblank"],
                    "zero_like_rows": counts[field]["zero_like"],
                    "nonzero_nonblank_rows": counts[field]["nonzero_nonblank"],
                    "legacy_mapping_status": mapping[0] if mapping else "UNBOUND",
                    "semantic_binding": "legacy_candidate_requires_current_validation" if mapping else "no_legacy_semantic_binding",
                    "legacy_mapping_reference": f"RDB_ONTOLOGY_MAPPING.md:{mapping[1]}" if mapping else "",
                    "unresolved_reason": (
                        "Legacy mapping is a read-only candidate; current unit/grain/as-of/coverage validation remains required"
                        if mapping
                        else "No matching legacy mapping row; retain raw schema meaning without semantic promotion"
                    ),
                }
            )
    return catalog, source_summary, all_fields


def profile_for_term(file_name: str, term: str, kinds: str) -> tuple[str, str, str, str]:
    """Return status, evidence code, unresolved reason, and deterministic profile name."""
    if term in {"bondkr:AverageAnnualTaxYield", "fpa:Return1Week", "etfkr:AverageMaturity", "etfkr:EffectiveDuration", "etfkr:ModifiedDuration"}:
        evidence = {
            "bondkr:AverageAnnualTaxYield": "PRBD01N001 avg_annual_tax_yield is zero in every row; stage_0c §4.3",
            "fpa:Return1Week": "PRFD01N001 fd_wk1_ern_r is entirely blank; stage_0c §4.4",
            "etfkr:AverageMaturity": "PREF01N001 fn_average_maturity is entirely blank; stage_0c §6",
            "etfkr:EffectiveDuration": "PREF01N001 fn_effective_duration is entirely blank; stage_0c §6",
            "etfkr:ModifiedDuration": "PREF01N001 fn_modified_duration is entirely blank; stage_0c §6",
        }[term]
        return ("제외", evidence, "No current official source value supports a semantic binding under the approved zero/missing handling", "empty_metric_binding")
    if "sh:NodeShape" in kinds:
        return (
            "수정",
            "TTL shape + CROSSWALK entity/relation audit",
            "Legacy shapes predate Portfolio/raw-holding/as-of revisions; revalidate after approved grain only",
            "shape_revalidation",
        )
    if term in {
        "fpa:Theme",
        "fpa:ThemeAssignment",
        "fpa:assignedEntity",
        "fpa:assignedTheme",
        "fpa:ControlRelationship",
        "fpa:parentOrganization",
        "fpa:subsidiaryOrganization",
        "fpa:controlType",
        "fpa:ownershipPercentage",
        "fpa:hasSubsidiary",
    }:
        return ("제외", "No verified theme/control source in current official or holdings inputs", "No current source binding; this does not prohibit future questions", "no_current_source")
    if term in {
        "fpa:BenchmarkIndex",
        "fpa:BenchmarkAssignment",
        "fpa:hasBenchmarkAssignment",
        "fpa:benchmarkedProduct",
        "fpa:assignedBenchmark",
        "fpa:assignmentKind",
        "fpa:Organization",
        "fpa:MarketParticipant",
        "fpa:Role",
        "fpa:ProductRole",
        "fpa:IssuerRole",
        "fpa:FundManagerRole",
        "fpa:AuthorizedParticipantRole",
        "fpa:LiquidityProviderRole",
        "fpa:BrokerRole",
        "fpa:TrusteeRole",
        "fpa:CustodianRole",
        "fpa:DistributorRole",
        "fpa:InvestorRole",
        "fpa:hasProductRole",
        "fpa:roleForProduct",
        "fpa:rolePlayedBy",
        "fpa:managedBy",
        "fpa:issuedByOrganization",
        "fpa:tracksBenchmark",
    }:
        return ("미확정", "Official source has raw organization/benchmark values; stage_0c has no verified entity resolution", "Name/code values cannot be merged into Organization or Benchmark entities without resolution evidence", "entity_resolution_pending")
    if term in {
        "fpa:PortfolioHoldingObservation",
        "fpa:portfolioProduct",
        "fpa:heldSecurity",
        "fpa:holdsSecurity",
        "fpa:weightValue",
        "fpa:quantityValue",
        "fpa:marketValue",
        "fpa:HoldingObservationShape",
    }:
        return ("수정", "holdings ZIP normalized mapping/holdings/coverage/failures + stage_0c §7", "Portfolio must be distinct from class/product; raw holdings can lack resolved Security; direct/look-through and as-of status are absent", "holdings_grain_revision")
    if term in {
        "fpa:ProjectProduct",
        "fpa:Security",
        "fpa:Listing",
        "fpa:FundSeries",
        "fpa:FundShareClass",
        "fpa:hasShareClass",
        "fpa:shareClassOf",
        "fpa:observedEntity",
        "fpa:observedProduct",
        "fpa:observedListing",
        "fpa:Observation",
        "fpa:ObservationContext",
        "fpa:MetricObservation",
        "fpa:MetricType",
        "fpa:valueStatus",
        "fpa:asOfDate",
        "fpa:periodStart",
        "fpa:periodEnd",
        "fpa:validFrom",
        "fpa:validTo",
        "fpa:collectedAt",
        "fpa:observationContext",
    }:
        return ("수정", "official field catalog + holdings bundle + stage_0c §4/§7", "Actual row/result grains and requested/effective as-of require a more specific binding", "core_grain_revision")
    if term in {
        "fpa:DataSource",
        "fpa:source",
        "fpa:sourceRecordId",
        "fpa:EvidenceDocument",
        "fpa:evidenceDocument",
        "fpa:joinMethod",
        "fpa:confidence",
        "fpa:conflictGroupId",
        "fpa:materializedShortcut",
    }:
        return ("유지", "MIGRATION_MANIFEST + holdings MANIFEST/raw records + architecture evidence boundary", "Concrete source record vocabulary remains valid; bindings are deferred", "provenance_kept")
    if term in {
        "fpa:ValueStatus",
        "fpa:ValidStatus",
        "fpa:ZeroExcludedStatus",
        "fpa:MissingStatus",
        "fpa:SentinelExcludedStatus",
        "fpa:ExternallyFilledStatus",
        "fpa:ConflictingStatus",
    }:
        return ("수정", "stage_0c §6-§7 and ARCHITECTURE §5", "Value status is valid but cannot carry independent coverage/failure/freshness states", "value_status_separation")
    if term in {
        "fpa:Identifier",
        "fpa:ISINIdentifier",
        "fpa:RICIdentifier",
        "fpa:TickerIdentifier",
        "fpa:ProjectIdentifier",
        "fpa:IdentifierScheme",
        "fpa:hasIdentifier",
        "fpa:identifies",
        "fpa:identifierScheme",
        "fpa:identifierValue",
        "fpa:verificationStatus",
        "fpa:ISINScheme",
        "fpa:RICScheme",
        "fpa:TickerScheme",
        "fpa:KSDItemNumberScheme",
        "fpa:ProjectProductNumberScheme",
    }:
        return ("수정", "official identifier candidates + holdings identifier observations; stage_0c §4/§7", "Schemes and verification status must be source-specific; placeholder/ambiguous/unresolved values are not identifiers", "identifier_revision")
    if term in {"fpa:MICScheme", "fpa:LEIScheme", "fpa:DARTCorpCodeScheme", "fpa:SECCIKScheme", "fpa:SECSeriesScheme", "fpa:SECClassScheme"}:
        return ("미확정", "Legacy mapping candidate only", "Current inputs contain candidates but no verified scheme mapping or crosswalk", "identifier_scheme_pending")
    if term in {"fpa:Exchange", "fpa:listedOnExchange", "fpa:securityOfListing", "fpa:hasListing", "fpa:tradingCurrency", "fpa:listingStartDate", "fpa:listingEndDate", "fpa:productCurrency", "fpa:Currency", "fpa:currency", "fpa:unitCode"}:
        return ("수정", "official field catalog and stage_0c §4.1-§4.4", "Source codes/currencies exist but exchange code schemes, unit semantics, and listing/product distinction are not uniformly verified", "listing_currency_revision")
    if term in {"fpa:SaleStatus", "fpa:SaleStatusObservation", "fpa:TradingStatus", "fpa:TradingStatusObservation", "fpa:tradingListing", "fpa:tradingStatus", "fpa:saleProduct", "fpa:seller", "fpa:saleStatus"}:
        return ("수정", "official field catalog; stage_0c §4.1-§4.4", "Status categories remain separate, but raw code/blank meaning and seller scope need source-specific treatment", "status_observation_revision")
    if term in {"fpa:ClassificationObservation", "fpa:classificationScheme", "fpa:classificationCode", "fpa:classificationLabel", "fpa:numericValue", "fpa:textValue", "fpa:metricType"}:
        return ("수정", "official field catalog + RDB_ONTOLOGY_MAPPING legacy candidates", "Raw classifications and metric values need field-level source, unit/period, and code-meaning validation", "raw_classification_revision")
    if "fpa:MetricType" in kinds:
        return ("수정", "official field catalog + stage_0c metric coverage", "Metric existence does not establish source-specific period/unit/currency/freshness compatibility", "generic_metric_revision")
    if file_name == "bond_kr.ttl":
        if term in {"bondkr:KoreanBond", "bondkr:BondLotContext", "bondkr:contextBond", "bondkr:tradeSegment", "bondkr:lotSequence"}:
            return ("유지", "PRBD01N001 source-row uniqueness; stage_0c §4.3", "Metric binding/freshness remains a separate catalog concern", "bond_grain_kept")
        return ("수정", "PRBD01N001 official field catalog + stage_0c §4.3", "Coverage differs by metric and code/rating meanings are not all confirmed", "bond_field_revision")
    if file_name == "etf_kr.ttl":
        if term in {"etfkr:DomesticExchangeTradedProduct", "etfkr:DomesticETF", "etfkr:DomesticETN"}:
            return ("유지", "PREF01N001 pd_grp_no distinguishes ETF/ETN; stage_0c §4.1", "Target-universe policy remains separate", "domestic_etp_kept")
        return ("수정", "PREF01N001 official field catalog + stage_0c §4.1", "Field-specific dates, coverage, and source meanings require revision", "domestic_etf_revision")
    if file_name == "etf_gl.ttl":
        if term in {"etfgl:SECRegisteredFundSeries", "etfgl:SECRegisteredShareClass", "etfgl:SECRegistrantRole", "etfgl:secFundSeries", "etfgl:secShareClass"}:
            return ("미확정", "PREF02N001 candidates and holdings raw SEC mapping", "Series/class/registrant resolution is not yet verified against the current source", "sec_resolution_pending")
        if term in {"etfgl:GlobalExchangeTradedProduct", "etfgl:GlobalETF", "etfgl:GlobalETN", "etfgl:fundInceptionDate"}:
            return ("수정", "PREF02N001 official field catalog + stage_0c §4.2", "ISIN/listing and stale row-date handling need current binding validation", "global_etf_revision")
        return ("수정", "PREF02N001 official field catalog + stage_0c §4.2", "Metric coverage and unit/freshness differ by field", "global_metric_revision")
    if file_name == "fund_pub.ttl":
        if term in {"fundpub:PublicFundShareClass", "fundpub:PubliclyOfferedFund"}:
            return ("유지", "PRFD01N001 itm_no class grain; stage_0c §4.4", "Portfolio is a separate additional grain", "fund_class_kept")
        if term in {"fundpub:RepresentativeFund", "fundpub:PublicFundShareClassShape", "fundpub:RepresentativeFundShape"}:
            return ("수정", "PRFD01N001 + holdings class→portfolio mapping; stage_0c §4.4/§7.3", "Representative grouping cannot be automatically identical to portfolio", "fund_portfolio_revision")
        return ("수정", "PRFD01N001 official field catalog + stage_0c §4.4", "Flag/code semantics, source dates, and metric coverage need source-specific treatment", "fund_field_revision")
    return ("미확정", "TTL declaration inventory", "No profile rule matched; retain term without semantic binding", "unclassified_term")


def ttl_inventory() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for path in sorted(LEGACY_ONTOLOGY.glob("*.ttl")):
        for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            term, separator, rest = line.partition(" a ")
            if not separator or ":" not in term or not all(character.isalnum() or character in "_:-" for character in term):
                continue
            kinds = [declaration for declaration in DECLARATION_TYPES if declaration in rest]
            if not kinds:
                continue
            status, evidence, unresolved, profile = profile_for_term(path.name, term, "+".join(kinds))
            rows.append(
                {
                    "ttl_file": path.name,
                    "ttl_sha256": sha256(path),
                    "line": str(line_number),
                    "term": term,
                    "declaration_types": "+".join(kinds),
                    "audit_status": status,
                    "evidence": evidence,
                    "unresolved_reason": unresolved,
                    "audit_profile": profile,
                }
            )
    return rows


def holdings_catalog() -> tuple[list[dict[str, object]], list[dict[str, object]], list[dict[str, object]], dict[str, int]]:
    members: list[dict[str, object]] = []
    fields: list[dict[str, object]] = []
    coverage_rows: list[dict[str, object]] = []
    with zipfile.ZipFile(HOLDINGS_ZIP) as archive, tempfile.TemporaryDirectory(prefix="canna_0d_") as temporary:
        temp_root = Path(temporary)
        manifest_member = "holdings_20260829/MANIFEST.json"
        bundle_manifest = json.loads(archive.read(manifest_member).decode("utf-8"))
        manifest_entries = {str(entry["path"]): entry for entry in bundle_manifest["files"]}
        zip_files = {info.filename: info for info in archive.infolist() if not info.is_dir()}
        manifest_mismatches = 0
        for relative_path, entry in manifest_entries.items():
            member_name = f"holdings_20260829/{relative_path}"
            info = zip_files.get(member_name)
            if info is None:
                manifest_mismatches += 1
                continue
            with archive.open(info) as source:
                actual_hash = hashlib.sha256(source.read()).hexdigest()
            if info.file_size != entry["bytes"] or actual_hash != entry["sha256"]:
                manifest_mismatches += 1
        if manifest_mismatches:
            raise ValueError(f"holdings internal MANIFEST mismatch count: {manifest_mismatches}")
        if set(zip_files).difference({manifest_member, *(f'holdings_20260829/{path}' for path in manifest_entries)}):
            raise ValueError("holdings ZIP contains a file not covered by the internal MANIFEST")
        for info in archive.infolist():
            if info.is_dir():
                continue
            with archive.open(info) as source:
                digest = hashlib.sha256(source.read()).hexdigest()
            members.append(
                {
                    "member_path": info.filename,
                    "bytes": info.file_size,
                    "sha256": digest,
                    "kind": "normalized" if "/normalized/" in info.filename else "raw_or_provenance",
                    "semantic_binding": "see normalized field records" if "/normalized/" in info.filename else "raw/provenance only; not a runtime binding",
                }
            )
        parquet_members = [info.filename for info in archive.infolist() if not info.is_dir() and info.filename.endswith(".parquet")]
        for member in parquet_members:
            output = temp_root / Path(member).name
            with archive.open(member) as source, output.open("wb") as destination:
                shutil.copyfileobj(source, destination)
            table_name = output.stem
            description = duckdb.connect(":memory:").execute("DESCRIBE SELECT * FROM read_parquet(?)", [str(output)]).fetchall()
            if table_name == "product_mapping":
                grain = "product/class → portfolio mapping record"
                binding = "candidate semantic mapping; exact identifier schemes remain unresolved"
            elif table_name == "holdings":
                grain = "portfolio × snapshot × raw holding record"
                binding = "raw holding observation; Security relation is conditional on identifier resolution"
            elif table_name == "coverage":
                grain = "family/product/class/portfolio/call coverage observation (source-specific)"
                binding = "dynamic coverage metadata, not Ontology TTL values"
            else:
                grain = "source collection failure record"
                binding = "dynamic failure metadata, not a value status or false relation"
            for column_name, column_type, nullable, *_ in description:
                fields.append(
                    {
                        "normalized_table": table_name,
                        "member_path": member,
                        "source_row_grain": grain,
                        "field": column_name,
                        "physical_type": column_type,
                        "nullable": nullable,
                        "semantic_binding": binding,
                        "coverage_failure_handling": "preserve source status; do not coerce to zero/false",
                    }
                )
            if table_name == "coverage":
                columns = [row[0] for row in description]
                quoted = ", ".join(f'"{column}"' for column in columns)
                data = duckdb.connect(":memory:").execute(f"SELECT {quoted} FROM read_parquet(?)", [str(output)]).fetchall()
                coverage_rows = [dict(zip(columns, row, strict=True)) for row in data]
    return members, fields, coverage_rows, {
        "holdings_manifest_file_rows": len(manifest_entries),
        "holdings_manifest_mismatches": manifest_mismatches,
    }


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        raise ValueError(f"Refusing to write empty catalog: {path.name}")
    with path.open("w", newline="", encoding="utf-8") as destination:
        writer = csv.DictWriter(destination, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as source:
        return list(csv.DictReader(source))


def verify_generated(output: Path) -> None:
    evidence_path = output / "audit_evidence.json"
    evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
    counts = evidence["counts"]
    terms = read_csv(output / "ttl_term_inventory.csv")
    official_fields = read_csv(output / "official_field_catalog.csv")
    official_sources = read_csv(output / "official_source_inventory.csv")
    members = read_csv(output / "holdings_bundle_member_inventory.csv")
    normalized_fields = read_csv(output / "holdings_normalized_field_catalog.csv")
    coverage = read_csv(output / "holdings_coverage_rows.csv")
    errors: list[str] = []
    if len(terms) != counts["ttl_term_rows"]:
        errors.append("ttl term row count differs from evidence")
    if any(not row["audit_status"].strip() or not row["evidence"].strip() or not row["unresolved_reason"].strip() for row in terms):
        errors.append("ttl term row has blank audit annotation")
    if len({(row["ttl_file"], row["term"]) for row in terms}) != len(terms):
        errors.append("ttl term key duplicate")
    if len(official_fields) != counts["official_field_rows"]:
        errors.append("official field row count differs from evidence")
    if any(not row["semantic_binding"].strip() for row in official_fields):
        errors.append("official field has blank semantic binding")
    if len({(row["table_id"], row["field"]) for row in official_fields}) != len(official_fields):
        errors.append("official field key duplicate")
    if len(official_sources) != counts["official_source_rows"]:
        errors.append("official source row count differs from evidence")
    if len(members) != counts["holdings_member_rows"] or len(normalized_fields) != counts["holdings_normalized_field_rows"] or len(coverage) != counts["holdings_coverage_rows"]:
        errors.append("holdings generated row count differs from evidence")
    for name, recorded in evidence["read_only_inputs"]["official_raw"].items():
        if sha256(OFFICIAL / name) != recorded:
            errors.append(f"official hash mismatch: {name}")
    if sha256(HOLDINGS_ZIP) != evidence["read_only_inputs"]["holdings_zip"]["sha256"]:
        errors.append("holdings ZIP hash mismatch")
    for name, recorded in evidence["read_only_inputs"]["legacy_ttl"].items():
        if sha256(LEGACY_ONTOLOGY / name) != recorded:
            errors.append(f"legacy TTL hash mismatch: {name}")
    if sha256(LEGACY_MAPPING) != evidence["read_only_inputs"]["legacy_mapping"]["sha256"]:
        errors.append("legacy mapping hash mismatch")
    _, _, _, manifest_check = holdings_catalog()
    if manifest_check != {
        "holdings_manifest_file_rows": counts["holdings_manifest_file_rows"],
        "holdings_manifest_mismatches": counts["holdings_manifest_mismatches"],
    }:
        errors.append("holdings internal MANIFEST verification differs from evidence")
    if errors:
        raise SystemExit("0-D verification failed: " + "; ".join(errors))
    print(
        json.dumps(
            {
                "status": "ok",
                "ttl_term_rows": len(terms),
                "official_field_rows": len(official_fields),
                "holdings_member_rows": len(members),
                "holdings_normalized_field_rows": len(normalized_fields),
                "holdings_manifest_mismatches": manifest_check["holdings_manifest_mismatches"],
            },
            ensure_ascii=False,
            sort_keys=True,
        )
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "provenance" / "workstreams" / "20260830_preintegration_parallel" / "ontology_data_alignment" / "generated",
    )
    parser.add_argument("--verify", action="store_true", help="verify existing generated catalogs without rewriting them")
    args = parser.parse_args()
    output = args.output.resolve()
    if args.verify:
        verify_generated(output)
        return
    output.mkdir(parents=True, exist_ok=True)

    official_fields, official_sources, _ = official_catalog()
    terms = ttl_inventory()
    holdings_members, holdings_fields, coverage, holdings_manifest_check = holdings_catalog()

    write_csv(output / "official_field_catalog.csv", official_fields)
    write_csv(output / "official_source_inventory.csv", official_sources)
    write_csv(output / "ttl_term_inventory.csv", terms)
    write_csv(output / "holdings_bundle_member_inventory.csv", holdings_members)
    write_csv(output / "holdings_normalized_field_catalog.csv", holdings_fields)
    write_csv(output / "holdings_coverage_rows.csv", coverage)

    evidence = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "generator": "scripts/audit_0d_inventory.py",
        "read_only_inputs": {
            "official_raw": {name: sha256(OFFICIAL / name) for table in TABLES.values() for name in (str(table["data"]), str(table["schema"]))},
            "holdings_zip": {"path": str(HOLDINGS_ZIP.relative_to(ROOT)), "sha256": sha256(HOLDINGS_ZIP)},
            "legacy_ttl": {path.name: sha256(path) for path in sorted(LEGACY_ONTOLOGY.glob("*.ttl"))},
            "legacy_mapping": {"path": str(LEGACY_MAPPING), "sha256": sha256(LEGACY_MAPPING)},
        },
        "counts": {
            "official_field_rows": len(official_fields),
            "official_source_rows": len(official_sources),
            "ttl_term_rows": len(terms),
            "holdings_member_rows": len(holdings_members),
            "holdings_normalized_field_rows": len(holdings_fields),
            "holdings_coverage_rows": len(coverage),
            **holdings_manifest_check,
        },
        "guarantees": [
            "All official schema fields are emitted once after data-header/schema-order validation.",
            "All named legacy declarations of the requested categories, plus ValueStatus/IdentifierScheme/AnnotationProperty declarations, are emitted once.",
            "All ZIP members are hash-inventoried and all normalized parquet fields are described from temporary copies.",
            "The script does not write data/official_raw or data/incoming and does not create a DB, Registry, Ontology, or runtime binding.",
        ],
    }
    (output / "audit_evidence.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(evidence["counts"], ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
